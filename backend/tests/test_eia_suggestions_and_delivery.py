from io import BytesIO
from types import SimpleNamespace as Obj
from unittest.mock import Mock, patch
from uuid import uuid4
import pytest
from docx import Document
from fastapi import HTTPException
from pydantic import ValidationError
from app.schemas.eia import EiaEmailRequest
from app.services.docx_content import append_editor_html
from app.services.docx_charts import content_volume_chart, subsection_distribution_chart
from app.services.eia_delivery_service import email_eia_document
from app.services.section_suggestions_service import generate_section_suggestions
from app.services.executive_summary_service import generate_executive_summary
from app.services.eia_document_export_service import build_compiled_eia_docx_bytes


def test_suggestions_are_non_mutating_and_source_labelled():
    sub = Obj(id=uuid4(), subsection_number='1.1', title='Objectives', content_html='', content='', checklist_mappings=[Obj(checklist_title='Explain objectives')], attachments=[], source_mappings=[])
    section = Obj(id=uuid4(), title='Project description', subsections=[sub])
    with patch('app.services.section_suggestions_service.get_eia_document_for_tenant', return_value=Obj(sections=[section])), patch('app.services.section_suggestions_service.get_settings', return_value=Obj(openai_api_key=None)):
        result = generate_section_suggestions(Mock(), Mock(), uuid4(), section.id)
    assert result['engine'] == 'checklist'
    assert len(result['sources']) == 2
    assert any('evidence' in p for p in result['subsections'][0]['suggestions'])
    assert 'draft' not in result['subsections'][0]
    assert sub.content_html == ''


def test_wrong_section_is_rejected():
    with patch('app.services.section_suggestions_service.get_eia_document_for_tenant', return_value=Obj(sections=[])), pytest.raises(HTTPException) as error:
        generate_section_suggestions(Mock(), Mock(), uuid4(), uuid4())
    assert error.value.status_code == 404


def test_docx_preserves_rich_text_and_tables():
    doc = Document()
    append_editor_html(doc, '<h2>Findings</h2><p>A <strong>significant</strong> impact.</p><ul><li>Mitigate dust</li></ul><table><tr><th>Impact</th><th>Measure</th></tr><tr><td>Dust</td><td>Water</td></tr></table>')
    buffer = BytesIO(); doc.save(buffer)
    restored = Document(BytesIO(buffer.getvalue()))
    assert any(run.bold and run.text == 'significant' for p in restored.paragraphs for run in p.runs)
    assert any(p.style.name == 'List Bullet' for p in restored.paragraphs)
    assert restored.tables[0].cell(1, 1).text == 'Water'


def test_email_attaches_exported_docx():
    with patch('app.services.eia_delivery_service.build_compiled_eia_docx_bytes', return_value=b'word-bytes'), patch('app.services.eia_delivery_service._send', return_value=True) as send:
        result = email_eia_document(Mock(), Obj(full_name='Author'), uuid4(), EiaEmailRequest(recipient='reader@example.com'))
    assert result['sent'] is True
    assert send.call_args.kwargs['attachments'][0][1] == b'word-bytes'
    assert send.call_args.kwargs['attachments'][0][0].endswith('.docx')


def test_export_access_failure_prevents_email():
    with patch('app.services.eia_delivery_service.build_compiled_eia_docx_bytes', side_effect=HTTPException(403)), patch('app.services.eia_delivery_service._send') as send, pytest.raises(HTTPException):
        email_eia_document(Mock(), Mock(), uuid4(), EiaEmailRequest(recipient='reader@example.com'))
    send.assert_not_called()


def test_email_failure_is_not_reported_as_sent():
    with patch('app.services.eia_delivery_service.build_compiled_eia_docx_bytes', return_value=b'word-bytes'), patch('app.services.eia_delivery_service._send', return_value=False), pytest.raises(HTTPException) as error:
        email_eia_document(Mock(), Obj(full_name='Author'), uuid4(), EiaEmailRequest(recipient='reader@example.com'))
    assert error.value.status_code == 502


@pytest.mark.parametrize('payload', [{'recipient': 'invalid'}, {'recipient': 'reader@example.com', 'subject': 'Report\r\nBcc: someone@example.com'}])
def test_email_input_validation(payload):
    with pytest.raises(ValidationError):
        EiaEmailRequest(**payload)


@pytest.mark.parametrize('fail', [False, True])
def test_ai_success_and_fallback_are_labelled(fail):
    import json
    sub = Obj(id=uuid4(), subsection_number='1.1', title='Objectives', content_html='<p>Draft</p>', content='', checklist_mappings=[], attachments=[], source_mappings=[])
    section = Obj(id=uuid4(), title='Project description', subsections=[sub])
    settings = Obj(openai_api_key='test', openai_evaluation_timeout_seconds=1, openai_evaluation_model='test')
    client = Mock()
    client.chat.completions.create.return_value = Obj(choices=[Obj(message=Obj(content=json.dumps({'summary': 'Add details', 'subsections': [{'subsection_id': str(sub.id), 'suggestions': ['Describe site access.']}]})))])
    if fail:
        client.chat.completions.create.side_effect = TimeoutError()
    with patch('app.services.section_suggestions_service.get_eia_document_for_tenant', return_value=Obj(sections=[section])), patch('app.services.section_suggestions_service.get_settings', return_value=settings), patch('openai.OpenAI', return_value=client):
        result = generate_section_suggestions(Mock(), Mock(), uuid4(), section.id)
    assert result['engine'] == ('checklist' if fail else 'ai')
    assert sub.content_html == '<p>Draft</p>'


def test_smtp_message_contains_docx_attachment():
    from app.services.email_service import _send
    with patch('app.services.email_service._smtp_connection', return_value=('smtp.example.test', 587, 'sender@example.test', None, None, True)), patch('app.services.email_service.smtplib.SMTP') as smtp:
        assert _send('Report', 'reader@example.test', 'Attached', attachments=[('report.docx', b'word', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')])
    message = smtp.return_value.__enter__.return_value.send_message.call_args.args[0]
    attachment = list(message.iter_attachments())[0]
    assert attachment.get_filename() == 'report.docx'
    assert attachment.get_payload(decode=True) == b'word'


def test_invalid_api_key_has_actionable_safe_message():
    class AuthenticationFailure(Exception):
        status_code = 401
    sub = Obj(id=uuid4(), subsection_number='1.1', title='Objectives', content_html='', content='', checklist_mappings=[], attachments=[], source_mappings=[])
    section = Obj(id=uuid4(), title='Project description', subsections=[sub])
    settings = Obj(openai_api_key='secret-test-key', openai_evaluation_timeout_seconds=1, openai_evaluation_model='test')
    client = Mock()
    client.chat.completions.create.side_effect = AuthenticationFailure('Do not expose secret-test-key')
    with patch('app.services.section_suggestions_service.get_eia_document_for_tenant', return_value=Obj(sections=[section])), patch('app.services.section_suggestions_service.get_settings', return_value=settings), patch('openai.OpenAI', return_value=client):
        result = generate_section_suggestions(Mock(), Mock(), uuid4(), section.id)
    assert 'AI authentication failed' in result['summary']
    assert 'secret-test-key' not in str(result)
    assert result['engine'] == 'checklist'


def test_report_charts_are_valid_png_images():
    from PIL import Image

    sections = [
        {
            'section_number': '1',
            'title': 'Project Description',
            'subsections': [{'content': 'A short project description with evidence.'}],
        },
        {
            'section_number': '2',
            'title': 'Alternatives',
            'subsections': [{'content': 'Options were compared.'}, {'content': 'The preferred option avoids impacts.'}],
        },
    ]
    for chart in (content_volume_chart(sections), subsection_distribution_chart(sections)):
        image = Image.open(chart)
        assert image.format == 'PNG'
        assert image.width >= 1000
        assert image.height >= 500


def test_ai_executive_summary_is_saved_to_document_metadata():
    import json

    subsection = Obj(subsection_number="1.1", title="Project", content_html="<p>Solar project.</p>", content="")
    section = Obj(section_number="1", title="Description", subsections=[subsection])
    project = Obj(name="Solar", country="Pakistan", location="Site", sector="Energy", description="Solar facility")
    document = Obj(title="Solar EIA", project=project, sections=[section], document_metadata={})
    database = Mock()
    response_data = {
        "project_overview": "The assessment covers a solar project.",
        "assessment_scope": "The saved EIA defines the assessment scope.",
        "baseline_conditions": "Baseline information is limited.",
        "key_effects": ["Construction effects require verification."],
        "alternatives": "The alternatives assessment should be confirmed.",
        "mitigation_and_monitoring": "Saved measures require assigned monitoring.",
        "conclusion_and_limitations": "The report contains evidence gaps.",
    }
    client = Mock()
    client.responses.create.return_value = Obj(output_text=json.dumps(response_data))
    settings = Obj(openai_api_key="test", openai_evaluation_timeout_seconds=1, openai_evaluation_model="test-model")

    with patch("app.services.executive_summary_service.get_eia_document_for_tenant", return_value=document), patch("app.services.executive_summary_service.get_settings", return_value=settings), patch("openai.OpenAI", return_value=client):
        result = generate_executive_summary(database, Mock(), uuid4())

    assert "The assessment covers" in result["html"]
    assert document.document_metadata["ai_executive_summary"]["model"] == "test-model"
    assert client.responses.create.call_args.kwargs["store"] is False
    database.commit.assert_called_once()


def test_docx_download_ensures_ai_summary_before_building_payload():
    payload = {
        "document": {"id": "doc", "title": "EIA", "status": "draft", "metadata": {}},
        "project": {"name": "Project", "country": None, "location": None, "sector": None, "description": None},
        "progress": {"progress_percentage": 0},
        "latest_evaluation": None,
        "sections": [],
    }
    with patch("app.services.eia_document_export_service.generate_executive_summary") as generate, patch("app.services.eia_document_export_service.build_compiled_eia_payload", return_value=payload):
        build_compiled_eia_docx_bytes(Mock(), Mock(), uuid4())
    generate.assert_called_once()
