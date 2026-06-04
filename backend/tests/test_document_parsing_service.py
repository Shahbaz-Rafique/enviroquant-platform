from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.services.document_parsing_service import parse_document_bytes


FIXTURE_TEXT = Path(__file__).with_name("fixtures").joinpath("sample_eia_excerpt.txt").read_text(encoding="utf-8")


def build_pdf_bytes(text: str) -> bytes:
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    y = 800
    for line in text.splitlines():
        pdf.drawString(40, y, line[:100])
        y -= 18
        if y < 60:
            pdf.showPage()
            y = 800
    pdf.save()
    return buffer.getvalue()


def build_docx_bytes(text: str) -> bytes:
    document = DocxDocument()
    for line in text.splitlines():
        if line.startswith("1.1.1") or line.startswith("4.6.1"):
            document.add_heading(line, level=1)
        elif line.strip():
            document.add_paragraph(line)
        else:
            document.add_paragraph("")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_parse_pdf_builds_traceable_chunks() -> None:
    parsed = parse_document_bytes(build_pdf_bytes(FIXTURE_TEXT), "sample-eia.pdf", "application/pdf")

    assert parsed.parser_status == "completed"
    assert parsed.metadata["chunk_count"] >= 1
    assert parsed.metadata["detected_sections"]
    assert any(chunk.page_number == 1 for chunk in parsed.chunks)
    assert any(chunk.section_number == "1.1.1" for chunk in parsed.chunks)


def test_parse_docx_builds_chunks_with_detected_sections() -> None:
    parsed = parse_document_bytes(
        build_docx_bytes(FIXTURE_TEXT),
        "sample-eia.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert parsed.parser_status == "completed"
    assert parsed.metadata["chunk_count"] >= 1
    assert parsed.metadata["detected_sections"]
    assert any(chunk.section_number in {"1.1.1", "4.6.1"} for chunk in parsed.chunks)


def test_parse_doc_marks_legacy_format_unsupported() -> None:
    parsed = parse_document_bytes(b"legacy-binary-placeholder", "legacy-report.doc", "application/msword")

    assert parsed.parser_status == "unsupported"
    assert parsed.chunks == []
    assert "not supported" in parsed.metadata["error"].lower()
