from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import re
from typing import Any


MAX_CHUNK_LENGTH = 1400
HEADING_PATTERN = re.compile(r"^\s*(\d+(?:\.\d+)+|\d+)\s+(.+?)\s*$")


@dataclass
class ParsedChunk:
    chunk_index: int
    page_number: int | None
    section_number: str | None
    section_title: str | None
    heading_path: list[str]
    content: str
    content_metadata: dict[str, Any]


@dataclass
class ParsedDocument:
    parser_status: str
    chunks: list[ParsedChunk]
    metadata: dict[str, Any]


def parse_document_bytes(content: bytes, filename: str, mime_type: str | None) -> ParsedDocument:
    extension = "." + filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if extension == ".pdf":
            return _parse_pdf(content)
        if extension == ".docx":
            return _parse_docx(content)
        if extension == ".doc":
            return ParsedDocument(
                parser_status="unsupported",
                chunks=[],
                metadata={
                    "error": "Legacy .doc parsing is not supported by the current parser pipeline.",
                    "detected_sections": [],
                    "chunk_count": 0,
                },
            )
        return ParsedDocument(
            parser_status="unsupported",
            chunks=[],
            metadata={"error": f"Unsupported parser extension: {extension}", "detected_sections": [], "chunk_count": 0},
        )
    except Exception as exc:
        return ParsedDocument(
            parser_status="failed",
            chunks=[],
            metadata={"error": str(exc), "detected_sections": [], "chunk_count": 0},
        )


def _parse_pdf(content: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    chunks: list[ParsedChunk] = []
    heading_path: list[str] = []
    current_section_number: str | None = None
    current_section_title: str | None = None

    for page_number, page in enumerate(reader.pages, start=1):
        raw_text = (page.extract_text() or "").strip()
        if not raw_text:
            continue
        page_chunks, current_section_number, current_section_title, heading_path = _extract_chunks_from_text(
            raw_text,
            page_number=page_number,
            starting_index=len(chunks),
            current_section_number=current_section_number,
            current_section_title=current_section_title,
            heading_path=heading_path,
        )
        chunks.extend(page_chunks)

    return ParsedDocument(
        parser_status="completed" if chunks else "empty",
        chunks=chunks,
        metadata=_parsed_metadata(chunks, parser_name="pypdf"),
    )


def _parse_docx(content: bytes) -> ParsedDocument:
    from docx import Document as DocxDocument

    document = DocxDocument(BytesIO(content))
    paragraphs: list[tuple[int | None, str, bool]] = []
    current_page = 1

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        has_page_break = _paragraph_has_page_break(paragraph)
        style_name = paragraph.style.name.lower() if paragraph.style and paragraph.style.name else ""
        is_heading_style = style_name.startswith("heading")
        if text:
            paragraphs.append((current_page, text, is_heading_style))
        if has_page_break:
            current_page += 1

    grouped_text: list[str] = []
    grouped_pages: list[int | None] = []
    for page_number, text, is_heading in paragraphs:
        prefix = "##HEADING## " if is_heading else ""
        grouped_text.append(prefix + text)
        grouped_pages.append(page_number)

    chunks: list[ParsedChunk] = []
    heading_path: list[str] = []
    current_section_number: str | None = None
    current_section_title: str | None = None
    buffer: list[str] = []
    buffer_pages: list[int | None] = []

    for page_number, raw_line in zip(grouped_pages, grouped_text, strict=False):
        forced_heading = raw_line.startswith("##HEADING## ")
        line = raw_line.replace("##HEADING## ", "", 1) if forced_heading else raw_line
        detected_heading = _detect_heading(line) if forced_heading or _detect_heading(line) else None
        if detected_heading:
            if buffer:
                chunks.append(
                    _build_chunk(
                        chunk_index=len(chunks),
                        page_number=_page_number_from_list(buffer_pages),
                        section_number=current_section_number,
                        section_title=current_section_title,
                        heading_path=heading_path,
                        text="\n\n".join(buffer),
                        parser_name="python-docx",
                        page_number_estimated=True,
                    )
                )
                buffer = []
                buffer_pages = []
            current_section_number, current_section_title = detected_heading
            heading_path = [f"{current_section_number} {current_section_title}".strip()]
            continue

        buffer.append(line)
        buffer_pages.append(page_number)
        if len("\n".join(buffer)) >= MAX_CHUNK_LENGTH:
            chunks.append(
                _build_chunk(
                    chunk_index=len(chunks),
                    page_number=_page_number_from_list(buffer_pages),
                    section_number=current_section_number,
                    section_title=current_section_title,
                    heading_path=heading_path,
                    text="\n\n".join(buffer),
                    parser_name="python-docx",
                    page_number_estimated=True,
                )
            )
            buffer = []
            buffer_pages = []

    if buffer:
        chunks.append(
            _build_chunk(
                chunk_index=len(chunks),
                page_number=_page_number_from_list(buffer_pages),
                section_number=current_section_number,
                section_title=current_section_title,
                heading_path=heading_path,
                text="\n\n".join(buffer),
                parser_name="python-docx",
                page_number_estimated=True,
            )
        )

    return ParsedDocument(
        parser_status="completed" if chunks else "empty",
        chunks=chunks,
        metadata=_parsed_metadata(chunks, parser_name="python-docx"),
    )


def _extract_chunks_from_text(
    text: str,
    *,
    page_number: int | None,
    starting_index: int,
    current_section_number: str | None,
    current_section_title: str | None,
    heading_path: list[str],
) -> tuple[list[ParsedChunk], str | None, str | None, list[str]]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    chunks: list[ParsedChunk] = []
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer, chunks
        if not buffer:
            return
        chunk_text = "\n".join(buffer).strip()
        if not chunk_text:
            buffer = []
            return
        for segment in _split_large_text(chunk_text):
            chunks.append(
                _build_chunk(
                    chunk_index=starting_index + len(chunks),
                    page_number=page_number,
                    section_number=current_section_number,
                    section_title=current_section_title,
                    heading_path=heading_path,
                    text=segment,
                    parser_name="pypdf",
                    page_number_estimated=False,
                )
            )
        buffer = []

    for line in lines:
        heading = _detect_heading(line)
        if heading:
            flush()
            current_section_number, current_section_title = heading
            heading_path = [f"{current_section_number} {current_section_title}".strip()]
            continue
        buffer.append(line)
        if len("\n".join(buffer)) >= MAX_CHUNK_LENGTH:
            flush()

    flush()
    return chunks, current_section_number, current_section_title, heading_path


def _detect_heading(line: str) -> tuple[str, str] | None:
    match = HEADING_PATTERN.match(line)
    if not match:
        return None
    section_number = match.group(1).strip()
    section_title = match.group(2).strip()
    if len(section_title) < 4:
        return None
    return section_number, section_title


def _split_large_text(text: str) -> list[str]:
    if len(text) <= MAX_CHUNK_LENGTH:
        return [text]
    paragraphs = [paragraph.strip() for paragraph in re.split(r"\n{2,}", text) if paragraph.strip()]
    segments: list[str] = []
    current: list[str] = []
    current_length = 0
    for paragraph in paragraphs:
        addition = len(paragraph) + 2
        if current and current_length + addition > MAX_CHUNK_LENGTH:
            segments.append("\n\n".join(current))
            current = [paragraph]
            current_length = len(paragraph)
            continue
        current.append(paragraph)
        current_length += addition
    if current:
        segments.append("\n\n".join(current))
    return segments


def _build_chunk(
    *,
    chunk_index: int,
    page_number: int | None,
    section_number: str | None,
    section_title: str | None,
    heading_path: list[str],
    text: str,
    parser_name: str,
    page_number_estimated: bool,
) -> ParsedChunk:
    compact = text.strip()
    return ParsedChunk(
        chunk_index=chunk_index,
        page_number=page_number,
        section_number=section_number,
        section_title=section_title,
        heading_path=heading_path,
        content=compact,
        content_metadata={
            "parser": parser_name,
            "page_number_estimated": page_number_estimated,
            "length": len(compact),
        },
    )


def _parsed_metadata(chunks: list[ParsedChunk], *, parser_name: str) -> dict[str, Any]:
    section_map: dict[str, dict[str, Any]] = {}
    for chunk in chunks:
        if not chunk.section_number:
            continue
        bucket = section_map.setdefault(
            chunk.section_number,
            {
                "section_number": chunk.section_number,
                "title": chunk.section_title or chunk.section_number,
                "content": "",
                "pages": [],
            },
        )
        if len(bucket["content"]) < 2400:
            bucket["content"] = (bucket["content"] + "\n\n" + chunk.content).strip()
        if chunk.page_number is not None and chunk.page_number not in bucket["pages"]:
            bucket["pages"].append(chunk.page_number)

    detected_sections = [
        {
            "section_number": item["section_number"],
            "title": item["title"],
            "content": item["content"][:2400],
            "metadata": {"pages": item["pages"]},
        }
        for item in section_map.values()
    ]
    return {
        "parser": parser_name,
        "chunk_count": len(chunks),
        "detected_sections": detected_sections,
    }


def _page_number_from_list(values: list[int | None]) -> int | None:
    numbers = [value for value in values if value is not None]
    return numbers[0] if numbers else None


def _paragraph_has_page_break(paragraph) -> bool:
    xml = paragraph._p.xml
    return 'w:type="page"' in xml
