from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
import textwrap
from uuid import UUID

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from sqlalchemy.orm import Session

from app.models.eia_evaluation import EiaEvaluationRun
from app.services.eia_evaluation_service import get_eia_evaluation_run
from app.services.eia_document_export_service import (
    _add_running_header_and_page_numbers,
    _apply_table_visual_system,
    _draw_pdf_header_footer,
    _pdf_chapter_banner,
    _pdf_safe_text,
    _pdf_table,
    _shade_docx_cell,
)


REPORT_HERO_PATH = Path(__file__).resolve().parent.parent / "assets" / "report" / "kuwait-environmental-intelligence-hero.png"


def _build_executive_dashboard_image(run) -> BytesIO:
    from PIL import Image, ImageDraw, ImageFont

    width, height = 1240, 1754
    canvas = Image.new("RGB", (width, height), "#F7F9F8")
    draw = ImageDraw.Draw(canvas)
    hero = Image.open(REPORT_HERO_PATH).convert("RGB")
    hero_ratio = max(width / hero.width, 500 / hero.height)
    hero = hero.resize((round(hero.width * hero_ratio), round(hero.height * hero_ratio)))
    hero_left = max(0, (hero.width - width) // 2)
    canvas.paste(hero.crop((hero_left, 0, hero_left + width, 500)), (0, 0))

    font_paths = (
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    )
    bold_paths = (
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    )

    def font(size: int, *, bold: bool = False):
        candidates = bold_paths if bold else font_paths
        for candidate in candidates:
            if Path(candidate).exists():
                return ImageFont.truetype(candidate, size=size)
        return ImageFont.load_default()

    def fit(text: object, max_chars: int) -> list[str]:
        return textwrap.wrap(str(text or ""), width=max_chars, break_long_words=False) or [""]

    def card(box, *, fill="#FFFFFF", outline="#D8E2DD", radius=10):
        draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=2)

    metadata = run.run_metadata or {}
    counts = metadata.get("status_counts") or {}
    findings = list(run.findings)
    total = len(findings) or 1
    missing = int(counts.get("MISSING", counts.get("MISSING_INFORMATION", 0)) or 0)
    compliant = int(counts.get("COMPLIANT", 0) or 0)
    partial = int(counts.get("PARTIALLY_COMPLIANT", 0) or 0)
    needs_work = int(counts.get("NEEDS_IMPROVEMENT", 0) or 0)
    needs_review = int(counts.get("NEEDS_REVIEW", 0) or 0)
    quality = round(float(metadata.get("overall_score") or 0) * 10)
    evidence_coverage = round(((total - missing) / total) * 100)
    confidence = round(sum(float(item.confidence_score or 0) for item in findings) / total * 100)
    risk = "High" if missing / total >= 0.45 else "Medium" if missing / total >= 0.2 else "Low"
    summary = str((metadata.get("review_report") or {}).get("summary") or "Review summary unavailable.")

    draw.text((60, 42), "ENVIROQUANT™", font=font(20, bold=True), fill="#A7E56F")
    draw.text((60, 112), "AI-Powered Environmental", font=font(42, bold=True), fill="white")
    draw.text((60, 166), "Impact Assessment", font=font(42, bold=True), fill="white")
    draw.text((60, 228), "Review & Compliance Report", font=font(29), fill="#8BD15F")
    for index, line in enumerate(fit(run.document.title, 58)[:2]):
        draw.text((60, 310 + index * 26), line, font=font(18, bold=index == 0), fill="#E4ECE8")
    draw.text((60, 390), "RQEIA assessment against Kuwait environmental requirements", font=font(16), fill="#C4D5CD")
    draw.text((60, 420), "Evidence-led findings • professional reviewer decision required", font=font(16), fill="#C4D5CD")

    metric_labels = ("OVERALL QUALITY", "EVIDENCE COVERAGE", "RISK LEVEL", "MISSING ITEMS", "AI CONFIDENCE")
    metric_values = (f"{quality}%", f"{evidence_coverage}%", risk, str(missing), f"{confidence}%")
    metric_notes = ("RQEIA index", "Requirements addressed", "Decision readiness", "Require attention", "Average confidence")
    gap = 12
    card_width = (1120 - gap * 4) // 5
    top = 535
    for index, (label, value, note) in enumerate(zip(metric_labels, metric_values, metric_notes)):
        left = 60 + index * (card_width + gap)
        card((left, top, left + card_width, top + 155))
        draw.text((left + 18, top + 18), label, font=font(13, bold=True), fill="#53675E")
        value_color = "#D58A16" if label == "RISK LEVEL" and risk != "Low" else "#236C4A"
        draw.text((left + 18, top + 62), value, font=font(34, bold=True), fill=value_color)
        draw.text((left + 18, top + 116), note, font=font(12, bold=True), fill="#60766C")

    draw.text((60, 730), "EXECUTIVE SUMMARY", font=font(24, bold=True), fill="#173C2E")
    for index, line in enumerate(fit(summary, 112)[:3]):
        draw.text((60, 770 + index * 25), line, font=font(16), fill="#40564C")

    strengths = [item.checklist_title for item in findings if item.status in {"COMPLIANT", "PARTIALLY_COMPLIANT"}][:3]
    gaps = [item.checklist_title for item in findings if item.status == "MISSING"][:3]
    actions = list((metadata.get("review_report") or {}).get("priority_actions") or [])[:2]
    panels = (
        ("STRENGTHS", strengths or ["Evidence-supported content recorded"], "#2E8B57"),
        ("GAPS", gaps or ["No material gap recorded"], "#D59A20"),
        ("RECOMMENDATIONS", actions or ["Complete priority actions and re-analyse"], "#D76A3A"),
        ("CONCLUSION", [summary], "#4F9A42"),
    )
    panel_gap = 14
    panel_width = (1120 - panel_gap * 3) // 4
    for index, (heading, items, accent) in enumerate(panels):
        left = 60 + index * (panel_width + panel_gap)
        card((left, 860, left + panel_width, 1115))
        draw.ellipse((left + 18, 880, left + 50, 912), outline=accent, width=3)
        draw.text((left + 62, 884), heading, font=font(14, bold=True), fill="#263F34")
        y = 930
        for item in items[:3]:
            lines = fit(item, 31)[:3]
            draw.text((left + 18, y), "•", font=font(15, bold=True), fill=accent)
            for line_index, line in enumerate(lines):
                draw.text((left + 34, y + line_index * 19), line, font=font(12), fill="#465B51")
            y += len(lines) * 19 + 12

    draw.text((60, 1160), "REGULATORY ALIGNMENT", font=font(23, bold=True), fill="#173C2E")
    regulatory = [item for item in findings if str(item.checklist_section).startswith("REG-KW-")]
    grouped: dict[str, list] = {}
    for item in regulatory:
        title = str((item.finding_metadata or {}).get("standard_title") or "Kuwait environmental requirements")
        grouped.setdefault(title, []).append(item)
    if not grouped:
        grouped = {"Kuwait environmental requirements": []}
    y = 1210
    for title, items in list(grouped.items())[:3]:
        score = round(sum(1 if item.status == "COMPLIANT" else 0.65 if item.status == "PARTIALLY_COMPLIANT" else 0.3 if item.status == "NEEDS_IMPROVEMENT" else 0 for item in items) / max(len(items), 1) * 100)
        draw.text((60, y), fit(title, 58)[0], font=font(14), fill="#344E43")
        draw.rounded_rectangle((60, y + 28, 500, y + 40), radius=6, fill="#E2E9E5")
        draw.rounded_rectangle((60, y + 28, 60 + round(440 * score / 100), y + 40), radius=6, fill="#4F9443")
        draw.text((520, y + 18), f"{score}%", font=font(16, bold=True), fill="#29463A")
        y += 72

    draw.text((650, 1160), "AI FINDINGS OVERVIEW", font=font(23, bold=True), fill="#173C2E")
    overview = (("COMPLIANT", compliant, "#438A47"), ("PARTIAL", partial + needs_work + needs_review, "#D79B1E"), ("MISSING", missing, "#D55345"), ("TOTAL", len(findings), "#3785B7"))
    for index, (label, value, accent) in enumerate(overview):
        left = 650 + (index % 2) * 265
        box_top = 1205 + (index // 2) * 95
        card((left, box_top, left + 245, box_top + 78))
        draw.text((left + 16, box_top + 13), str(value), font=font(27, bold=True), fill=accent)
        draw.text((left + 75, box_top + 23), label, font=font(13, bold=True), fill="#42584E")

    draw.text((60, 1455), "RECENT KEY FINDINGS", font=font(23, bold=True), fill="#173C2E")
    columns = (60, 610, 830, 1010)
    headers = ("FINDING", "REGULATION / AREA", "STATUS", "AI CONFIDENCE")
    draw.rectangle((60, 1495, 1180, 1535), fill="#EEF2F0")
    for x, label in zip(columns, headers):
        draw.text((x + 10, 1508), label, font=font(12, bold=True), fill="#41574D")
    key_findings = sorted(findings, key=lambda item: (item.status != "MISSING", -float(item.confidence_score or 0)))[:4]
    row_y = 1535
    for finding in key_findings:
        draw.line((60, row_y + 50, 1180, row_y + 50), fill="#DDE5E1", width=1)
        draw.text((70, row_y + 16), fit(finding.checklist_title, 64)[0], font=font(12), fill="#30473D")
        area = (finding.finding_metadata or {}).get("standard_code") or finding.checklist_section
        draw.text((620, row_y + 16), str(area)[:25], font=font(12), fill="#30473D")
        draw.text((840, row_y + 16), finding.status.replace("_", " ").title(), font=font(12, bold=True), fill="#B44E3E" if finding.status == "MISSING" else "#446F50")
        draw.text((1035, row_y + 16), f"{round(float(finding.confidence_score or 0) * 100)}%", font=font(12, bold=True), fill="#2E694B")
        row_y += 50

    output = BytesIO()
    canvas.save(output, format="PNG", optimize=True)
    output.seek(0)
    return output


def _create_docx_document():
    try:
        from docx import Document as create_document
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "DOCX report generation requires python-docx to be installed in the backend environment."
        ) from exc

    return create_document()


def build_report_payload(db: Session, current_user, document_id: UUID, run_id: UUID) -> dict:
    run = get_eia_evaluation_run(db, current_user, document_id, run_id)
    return {
        "run": {
            "id": str(run.id),
            "status": run.status,
            "prompt_version": run.prompt_version,
            "model_version": run.model_version,
            "overall_score": run.run_metadata.get("overall_score"),
            "overall_appraisal": run.run_metadata.get("overall_appraisal"),
        },
        "section_summaries": [
            {
                "section_number": summary.section_number,
                "section_title": summary.section_title,
                "score": summary.score,
                "summary_comment": summary.summary_comment,
            }
            for summary in run.section_summaries
        ],
        "findings": [
            {
                "checklist_section": finding.checklist_section,
                "checklist_title": finding.checklist_title,
                "status": finding.status,
                "adequacy": finding.adequacy,
                "confidence_score": finding.confidence_score,
                "evidence_summary": finding.evidence_summary,
                "ai_analysis": finding.ai_analysis,
                "recommendation": finding.recommendation,
                "missing_elements": finding.missing_elements,
                "evidence_references": finding.evidence_references,
                "reviewer_decision": (finding.finding_metadata or {}).get("reviewer_decision"),
                "reviewer_decision_note": (finding.finding_metadata or {}).get("reviewer_decision_note"),
            }
            for finding in run.findings
        ],
    }


def build_report_json_bytes(db: Session, current_user, document_id: UUID, run_id: UUID) -> bytes:
    payload = build_report_payload(db, current_user, document_id, run_id)
    return json.dumps(payload, indent=2).encode("utf-8")


def build_report_docx_bytes(db: Session, current_user, document_id: UUID, run_id: UUID) -> bytes:
    run = get_eia_evaluation_run(db, current_user, document_id, run_id)
    document = _create_docx_document()
    _setup_review_docx(document, run.document.title)
    _add_review_cover(document, run)
    document.add_page_break()

    document.add_heading("Document Control", level=1)
    control = document.add_table(rows=0, cols=2)
    control.style = "Table Grid"
    for label, value in (
        ("Reviewed EIA", run.document.title),
        ("Review run", str(run.id)),
        ("Methodology", getattr(run, "methodology_version", "RQEIA")),
        ("Checklist", getattr(run, "checklist_version", "RQEIA detailed questions")),
        ("Model", run.model_version),
        ("Completed", run.completed_at.isoformat() if run.completed_at else "Not completed"),
        ("Issue purpose", "Professional review and decision-readiness assessment"),
    ):
        cells = control.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)
        cells[0].paragraphs[0].runs[0].bold = True

    metadata = run.run_metadata or {}
    status_counts = metadata.get("status_counts") or {}
    document.add_heading("Executive Review Summary", level=1)
    document.add_paragraph(str((metadata.get("review_report") or {}).get("summary") or "Review summary unavailable."))
    metrics = document.add_table(rows=2, cols=5)
    metrics.style = "Table Grid"
    metric_values = (
        ("Quality index", f"{round(float(metadata.get('overall_score') or 0) * 10)}%"),
        ("Compliant", status_counts.get("COMPLIANT", 0)),
        ("Partial", status_counts.get("PARTIALLY_COMPLIANT", 0)),
        ("Needs work", status_counts.get("NEEDS_IMPROVEMENT", 0)),
        ("Missing", status_counts.get("MISSING", status_counts.get("MISSING_INFORMATION", 0))),
    )
    for index, (label, value) in enumerate(metric_values):
        metrics.rows[0].cells[index].text = label
        metrics.rows[0].cells[index].paragraphs[0].runs[0].bold = True
        metrics.rows[1].cells[index].text = str(value)

    priority_actions = (metadata.get("review_report") or {}).get("priority_actions") or []
    document.add_heading("Priority Actions", level=2)
    if priority_actions:
        for action in priority_actions:
            document.add_paragraph(str(action), style="List Bullet")
    else:
        document.add_paragraph("No priority actions were recorded in this run.")

    document.add_heading("RQEIA Area Assessment", level=1)
    for summary in run.section_summaries:
        document.add_heading(f"{summary.section_number}. {summary.section_title}", level=2)
        document.add_paragraph(
            f"Coverage index: {round(summary.score * 10)}% | Questions: {summary.findings_count} | "
            f"Compliant: {summary.compliant_count} | Partial: {summary.partially_compliant_count} | "
            f"Needs work: {summary.needs_improvement_count} | Missing: {summary.missing_count}"
        )
        if summary.summary_comment:
            document.add_paragraph(summary.summary_comment)

    document.add_page_break()
    document.add_heading("Detailed Findings and Evidence Traceability", level=1)
    for finding in run.findings:
        document.add_heading(f"{finding.checklist_section} {finding.checklist_title}", level=2)
        overview = document.add_table(rows=2, cols=3)
        overview.style = "Table Grid"
        for index, label in enumerate(("Assessment status", "Adequacy", "AI confidence")):
            overview.rows[0].cells[index].text = label
            overview.rows[0].cells[index].paragraphs[0].runs[0].bold = True
        overview.rows[1].cells[0].text = finding.status.replace("_", " ").title()
        overview.rows[1].cells[1].text = finding.adequacy.replace("_", " ").title()
        overview.rows[1].cells[2].text = f"{round(finding.confidence_score * 100)}%"
        document.add_paragraph(f"Evidence: {finding.evidence_summary}")
        document.add_paragraph(f"Assessment: {finding.ai_analysis}")
        if finding.recommendation:
            document.add_paragraph(f"Recommended action: {finding.recommendation}")
        if finding.missing_elements:
            document.add_paragraph("Gap / missing information:")
            for item in finding.missing_elements:
                document.add_paragraph(item, style="List Bullet")
        if finding.evidence_references:
            document.add_paragraph("Source references:")
            for reference in finding.evidence_references[:8]:
                filename = reference.get("source_document_filename") or "Structured EIA draft"
                page = f", page {reference['page_number']}" if reference.get("page_number") else ""
                excerpt = str(reference.get("excerpt") or "").strip()
                document.add_paragraph(f"{filename}{page}: {excerpt}", style="List Bullet")
        decision = (finding.finding_metadata or {}).get("reviewer_decision")
        document.add_paragraph(
            f"Reviewer decision: {str(decision).replace('_', ' ').title() if decision else 'Pending'}"
            + (f" — {(finding.finding_metadata or {}).get('reviewer_decision_note')}" if (finding.finding_metadata or {}).get("reviewer_decision_note") else "")
        )

    document.add_heading("Reviewer Determination", level=1)
    document.add_paragraph(
        "This AI-assisted review supports professional assessment and does not constitute regulatory approval. "
        "The authorised reviewer must confirm applicability, evidence sufficiency, appraisal and finding closure."
    )
    approval = document.add_table(rows=4, cols=4)
    approval.style = "Table Grid"
    for index, label in enumerate(("Role", "Name", "Decision / signature", "Date")):
        approval.rows[0].cells[index].text = label
        approval.rows[0].cells[index].paragraphs[0].runs[0].bold = True
    for row, role in zip(approval.rows[1:], ("EIA Lead", "Independent Reviewer", "Project Director")):
        row.cells[0].text = role

    _apply_table_visual_system(document)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _setup_review_docx(document, title: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10.5)
    document.styles["Normal"].font.color.rgb = RGBColor(0x18, 0x37, 0x2C)
    for level in (1, 2):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.bold = True
        if level == 1:
            style.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "123F2E")
            style.element.get_or_add_pPr().append(shading)
        else:
            style.font.color.rgb = RGBColor(0x28, 0x74, 0x51)
    _add_running_header_and_page_numbers(document.sections[0], title)


def _add_review_cover(document, run) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = 0
    paragraph.add_run().add_picture(_build_executive_dashboard_image(run), width=Inches(6.5))


def build_report_pdf_bytes(db: Session, current_user, document_id: UUID, run_id: UUID) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, PageBreak

    run = get_eia_evaluation_run(db, current_user, document_id, run_id)
    buffer = BytesIO()
    styles = getSampleStyleSheet()
    dark, green, lime, pale, ink, muted = (colors.HexColor(value) for value in ("#062A22", "#287451", "#8BD15F", "#F1F7F3", "#18372C", "#52675E"))
    body = ParagraphStyle("Review Body", parent=styles["BodyText"], fontSize=8.8, leading=12, textColor=ink, spaceAfter=5)
    small = ParagraphStyle("Review Small", parent=body, fontSize=7.2, leading=9, textColor=muted)
    chapter = ParagraphStyle("Review Chapter", parent=styles["Heading1"], textColor=colors.white, fontName="Helvetica-Bold", fontSize=15)
    subheading = ParagraphStyle("Review Subheading", parent=styles["Heading2"], textColor=green, fontName="Helvetica-Bold", fontSize=11, spaceBefore=8, spaceAfter=5)
    dashboard = Image(_build_executive_dashboard_image(run), width=174 * mm, height=246 * mm)
    story = [dashboard, PageBreak(), _pdf_chapter_banner("EXECUTIVE REVIEW SUMMARY", "RQEIA quality, evidence and decision-readiness position", chapter, small, dark, pale)]
    metadata = run.run_metadata or {}
    counts = metadata.get("status_counts") or {}
    metrics = [["QUALITY INDEX", "COMPLIANT", "PARTIAL", "NEEDS WORK", "MISSING"], [f"{round(float(metadata.get('overall_score') or 0) * 10)}%", counts.get("COMPLIANT", 0), counts.get("PARTIALLY_COMPLIANT", 0), counts.get("NEEDS_IMPROVEMENT", 0), counts.get("MISSING", counts.get("MISSING_INFORMATION", 0))]]
    story.extend([Spacer(1, 5 * mm), _pdf_table(metrics, [34 * mm] * 5, body, green, pale), Spacer(1, 5 * mm), Paragraph(_pdf_safe_text(str((metadata.get("review_report") or {}).get("summary") or "Review summary unavailable.")), body), Spacer(1, 5 * mm), _pdf_chapter_banner("RQEIA AREA ASSESSMENT", "Detailed assessment across all eight methodology areas", chapter, small, dark, pale)])
    for summary in run.section_summaries:
        story.extend([Paragraph(f"{summary.section_number}  {summary.section_title}", subheading), _pdf_table([["Coverage", "Questions", "Compliant", "Partial", "Needs work", "Missing"], [f"{round(summary.score * 10)}%", summary.findings_count, summary.compliant_count, summary.partially_compliant_count, summary.needs_improvement_count, summary.missing_count]], [28 * mm] * 6, small, green, pale), Paragraph(_pdf_safe_text(summary.summary_comment or "No area summary recorded."), body)])
    story.extend([PageBreak(), _pdf_chapter_banner("DETAILED FINDINGS & EVIDENCE TRACEABILITY", "Requirement → Evidence → Assessment → Finding → Recommended action → Reviewer decision", chapter, small, dark, pale)])
    for finding in run.findings:
        decision = (finding.finding_metadata or {}).get("reviewer_decision") or "PENDING"
        story.extend([Paragraph(f"{finding.checklist_section}  {finding.checklist_title}", subheading), _pdf_table([["Status", "Adequacy", "Confidence", "Reviewer decision"], [finding.status.replace("_", " ").title(), finding.adequacy.replace("_", " ").title(), f"{round(finding.confidence_score * 100)}%", str(decision).replace("_", " ").title()]], [42.5 * mm] * 4, small, green, pale), Paragraph(f"<b>Evidence:</b> {_pdf_safe_text(finding.evidence_summary)}", body), Paragraph(f"<b>Assessment:</b> {_pdf_safe_text(finding.ai_analysis)}", body), Paragraph(f"<b>Recommended action:</b> {_pdf_safe_text(finding.recommendation or 'No action recorded')}", body)])
        if finding.evidence_references:
            refs = []
            for reference in finding.evidence_references[:8]:
                source = reference.get("source_document_filename") or reference.get("subsection_number") or "Structured EIA"
                page = f", page {reference['page_number']}" if reference.get("page_number") else ""
                refs.append(f"{source}{page}")
            story.append(Paragraph(f"<b>Sources:</b> {_pdf_safe_text('; '.join(refs))}", small))
    story.extend([PageBreak(), _pdf_chapter_banner("REVIEWER DETERMINATION", "Final approval remains with the authorised human reviewer", chapter, small, dark, pale), Paragraph("This AI-assisted review does not constitute regulatory approval. Findings may only be closed or retained by an authorised reviewer after examining the revised content and supporting evidence.", body), _pdf_table([["Role", "Name", "Decision / signature", "Date"], ["EIA Lead", "", "", ""], ["Independent Reviewer", "", "", ""], ["Project Director", "", "", ""]], [42.5 * mm] * 4, body, dark, pale)])
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=19 * mm, bottomMargin=18 * mm, title="EIA Quality Review & Decision Readiness Report", author="EnviroQuant")
    doc.build(story, onFirstPage=lambda canvas, doc: None, onLaterPages=lambda canvas, doc: _draw_pdf_header_footer(canvas, doc, run.document.title))
    return buffer.getvalue()
