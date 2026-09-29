"""Convert supported editor HTML into Word paragraphs and tables without fetching URLs."""
from html.parser import HTMLParser


class DocxContentParser(HTMLParser):
    def __init__(self, document):
        super().__init__(convert_charrefs=True)
        self.document = document
        self.paragraph = None
        self.bold = False
        self.italic = False
        self.lists = []
        self.table = None
        self.cell = None
        self.column = 0
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'):
            self.skip += 1
        if self.skip:
            return
        if tag in ('ul', 'ol'):
            self.lists.append(tag)
        elif tag == 'table':
            self.table = self.document.add_table(rows=0, cols=1)
            self.table.style = 'Table Grid'
        elif tag == 'tr' and self.table is not None:
            self.table.add_row()
            self.column = 0
        elif tag in ('td', 'th') and self.table is not None:
            from docx.shared import Inches
            if not self.table.rows:
                self.table.add_row()
            if self.column >= len(self.table.columns):
                self.table.add_column(Inches(1.2))
            self.cell = self.table.rows[-1].cells[self.column]
            self.column += 1
            self.paragraph = self.cell.paragraphs[0]
            self.bold = tag == 'th'
        elif tag in ('p', 'h1', 'h2', 'h3', 'li', 'blockquote'):
            container = self.cell if self.cell is not None else self.document
            style = ('List Number' if self.lists and self.lists[-1] == 'ol' else 'List Bullet') if tag == 'li' else None
            if tag.startswith('h') and len(tag) == 2:
                style = f'Heading {min(int(tag[1]) + 2, 5)}'
            self.paragraph = container.add_paragraph(style=style)
        elif tag in ('strong', 'b'):
            self.bold = True
        elif tag in ('em', 'i'):
            self.italic = True
        elif tag == 'br' and self.paragraph is not None:
            self.paragraph.add_run().add_break()
        elif tag == 'img':
            values = dict(attrs)
            self.handle_data('[Figure: ' + (values.get('alt') or 'See source image') + ']')
            if values.get('src', '').startswith(('https://', 'http://')):
                self.handle_data(' ' + values['src'])

    def handle_endtag(self, tag):
        if tag in ('script', 'style'):
            self.skip = max(0, self.skip - 1)
        if tag in ('strong', 'b', 'th'):
            self.bold = False
        if tag in ('em', 'i'):
            self.italic = False
        if tag in ('ul', 'ol') and self.lists:
            self.lists.pop()
        if tag in ('p', 'h1', 'h2', 'h3', 'li', 'blockquote', 'td', 'th', 'table'):
            self.paragraph = None
        if tag in ('td', 'th'):
            self.cell = None
        if tag == 'table':
            self.table = None

    def handle_data(self, data):
        if self.skip or (not data.strip() and self.paragraph is None):
            return
        if self.paragraph is None:
            self.paragraph = self.document.add_paragraph()
        run = self.paragraph.add_run(data)
        run.bold = self.bold
        run.italic = self.italic


def append_editor_html(document, html):
    parser = DocxContentParser(document)
    parser.feed(html)
    parser.close()
